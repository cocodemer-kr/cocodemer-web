#!/usr/bin/env python3
"""Turn the original active-ingredient workbook into ingredients.json.

Source: 전산개발/홈페이지/맞춤형 액티브 문안_최종-20210208.xlsx — one sheet per
active, laid out as label row / value row. Two codes appear twice in the
original (Aa, Sl), so sheets are keyed by sheet name and given a unique id.
"""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, '_i18n', 'ingredients.json')

GROUP_SHEETS = {'미백액티브': '미백', '주름액티브': '주름', '보습액티브': '보습',
                '민감성액티브': '민감성', '베이스': '베이스'}
LABELS = ['원료소개', '효능요약', '원료 상세내용', '더 알아보기', '연관 처방제품']

# Names are taken from the workbook's own efficacy line where it gives one, and
# otherwise from the ingredient named in the introduction — cross-checked
# against the ACTIVE CODE lines already printed on the product pages.
# Three actives are named nowhere in the source; those are left blank rather
# than guessed, and the page says so.
NAMES = {
    'NP': ('', ''), 'Pt': ('', ''), 'Ma': ('', ''),
    'Ds': ('디메칠설폰', 'Dimethyl Sulfone'),
    'Hp': ('헤스페리딘', 'Hesperidin'),
    'Ta': ('트라넥삼산', 'Tranexamic Acid'),
    'At': ('아스코빌테트라이소팔미테이트', 'Ascorbyl Tetraisopalmitate'),
    'Cs': ('카르노신', 'Carnosine'),
    'Na': ('나이아신아마이드', 'Niacinamide'),
    'Ah': ('아세틸헥사펩타이드-8', 'Acetyl Hexapeptide-8'),
    'Pp': ('팔미토일펜타펩타이드-4', 'Palmitoyl Pentapeptide-4'),
    'Ka': ('타라열매·코토니 추출물 복합체', 'FILMEXEL® by SILAB'),
    'So': ('에스에이치올리고펩타이드-1', 'sh-Oligopeptide-1'),
    'Cn': ('세라마이드엔피', 'Ceramide NP'),
    'Ad': ('아데노신', 'Adenosine'),
    'Ui': ('유비퀴논', 'Ubiquinone'),
    'Pm': ('팔미토일트라이펩타이드-38', 'Palmitoyl Tripeptide-38'),
    'Sd': ('소듐디엔에이', 'Sodium DNA'),
    'Sh': ('소듐하이알루로네이트', 'Sodium Hyaluronate'),
    'Rf': ('라피노스', 'Raffinose'),
    'Ct': ('크레아틴', 'Creatine'),
    'Hh': ('하이드록시프로필트라이모늄하이알루로네이트',
           'Hydroxypropyltrimonium Hyaluronate'),
    'Is': ('이노시톨', 'Inositol'),
    'Bt': ('베타인', 'Betaine'),
    'Pa': ('폴리글루타믹애씨드', 'Polyglutamic Acid'),
    'Ab': ('알로에베라잎추출물', 'Aloe Barbadensis Leaf Extract'),
    'Th': ('트레할로오스', 'Trehalose'),
    'Pn': ('판테놀', 'Panthenol'),
    'Ca': ('병풀추출물', 'Centella Asiatica Extract'),
    'Ar': ('지모뿌리추출물', 'Anemarrhena Asphodeloides Root Extract'),
    'Aa (2)': ('개똥쑥추출물', 'Artemisia Annua Extract'),
    'Ac': ('아세틸다이펩타이드-1세틸에스터', 'Acetyl Dipeptide-1 Cetyl Ester'),
    'It': ('이소트레티노인', 'Isotretinoin'),
    'As': ('아시아티코사이드', 'Asiaticoside'),
    'Bf': ('바이오플라보노이드', 'Bioflavonoids'),
    'Ms': ('메도우폼씨오일', 'Limnanthes Alba (Meadowfoam) Seed Oil'),
    'Sb': ('시어버터', 'Shea Butter'),
    'HI': ('하이드로제네이티드레시틴', 'Hydrogenated Lecithin'),
    'Sl (2)': ('스피루리나', 'Spirulina'),
    'Pl': ('폴리글리세릴계 유화제', ''),
    'Ao': ('아르간오일', 'Argan Oil'),
    'Co': ('올리브 유래 유화왁스', ''),
    'Sl': ('스쿠알란', 'Squalane'),
    'Tt': ('토코트리에놀', 'Tocotrienols'),
    'Aa': ('아시아틱애씨드', 'Asiatic Acid'),
}


# Two codes appear twice in the original workbook. Pin the ids so links from
# the product pages land on the right one: a product's "Sl" is squalane and its
# "Aa" is asiatic acid, so those keep the bare code.
ID_OVERRIDE = {'Sl (2)': 'Sl-spirulina', 'Aa (2)': 'Aa-artemisia'}

# Two entries are a class of emulsifier rather than a single INCI name, so they
# get a plain English description instead of an INCI.
NAME_EN = {'Co': 'Olive-Derived Emulsifying Wax', 'Pl': 'Polyglyceryl Emulsifier'}


def sheet_fields(sheet):
    rows = [['' if v is None else str(v).strip() for v in r]
            for r in sheet.iter_rows(values_only=True)]
    found = {'code': ''}
    i = 0
    while i < len(rows):
        label = next((c for c in rows[i] if c in LABELS), None)
        if label:
            j = i + 1
            while j < len(rows) and not any(rows[j]):
                j += 1
            if j < len(rows):
                values = [c for c in rows[j] if c]
                if label == '원료소개' and len(rows[j]) > 2 and rows[j][2]:
                    found['code'] = rows[j][2]
                if values:
                    found[label] = max(values, key=len)
                    urls = [c for c in values if c.startswith('http')]
                    if urls:
                        found['url'] = urls[0]
            i = j
        i += 1
    return found


def bullets(text):
    return [l.strip(' -–—·') for l in (text or '').split('\n') if l.strip(' -–—·')]


def parse(path):
    book = openpyxl.load_workbook(path, data_only=True)
    out, group, used = [], None, {}
    for name in book.sheetnames:
        if name in GROUP_SHEETS:
            group = GROUP_SHEETS[name]
            continue
        f = sheet_fields(book[name])
        code = (f['code'] or re.sub(r'\s*\(\d\)$', '', name)).strip()
        if code.upper() == 'HI':
            code = 'Hl'
        ko, inci = NAMES.get(name, ('', ''))

        efficacy = f.get('효능요약', '')
        detail = f.get('원료 상세내용', '')
        if not efficacy and detail and not detail.startswith('http'):
            efficacy = detail          # one sheet files its benefits here
        lines = efficacy.split('\n')
        if ko and lines and lines[0].replace(' ', '').startswith(ko.replace(' ', '')[:4]):
            efficacy = '\n'.join(lines[1:])

        uid = ID_OVERRIDE.get(name)
        if not uid:
            uid = code if code not in used else f'{code}-{used[code] + 1}'
            used[code] = used.get(code, 0) + 1
        out.append({
            'id': uid, 'code': code, 'group': group,
            'name_ko': ko, 'inci': inci, 'name_en': NAME_EN.get(name, ''),
            'intro': f.get('원료소개', ''),
            'benefits': bullets(efficacy),
            'deep': f.get('더 알아보기', ''),
            'source': f.get('url', ''),
            'pending_name': not ko,
        })
    return out


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else (
        '/mnt/user-data/uploads/COCODEMER/전산개발/홈페이지/'
        '맞춤형 액티브 문안_최종-20210208.xlsx')
    rows = parse(src)
    json.dump(rows, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'  {len(rows)} actives -> {OUT}')
    print('  이름 미확정:', [r['code'] for r in rows if r['pending_name']])
    print('  효능 없음:', [r['code'] for r in rows if not r['benefits']])
    print('  소개 없음:', [r['code'] for r in rows if not r['intro']])
